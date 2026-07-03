# -*- coding = utf-8 -*-
# @Time :2024/6/1
# @Author :Claude Code
# @File :  export_docx.py
"""
学习通作业导出为Word(docx)模块
支持将JSON格式的答案文件导出为格式化的Word文档
"""
import os
import re
from datetime import datetime
from json import JSONDecodeError

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from my_xxt.answer_files import answer_json_files, read_answer_json


TYPE_COLORS = {
    "单选题": (0x40, 0x9E, 0xFF),
    "多选题": (0x67, 0xC2, 0x3A),
    "判断题": (0xE6, 0xA2, 0x3C),
    "填空题": (0x90, 0x94, 0x99),
    "简答题": (0x00, 0xB5, 0x78),
    "论述题": (0x9C, 0x27, 0xB0),
    "编程题": (0xF5, 0x6C, 0x6C),
    "其他": (0x60, 0x60, 0x60),
}
ANSWER_COLOR = (0xC0, 0x44, 0x32)
FOOTER_TEXT = "— 本文件由 new_xxt 学习通助手自动生成 —"


def sanitize_filename(filename: str, replacement: str = "_") -> str:
    """Return a filename safe for Windows, macOS, and Linux."""
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', replacement, filename).strip()
    cleaned = cleaned.rstrip(". ")
    return cleaned or "未命名"


def course_matches_query(course_name: str, query: str) -> bool:
    query = query.strip()
    if not query:
        return False
    return query in course_name


def set_cell_shading(cell, color):
    """设置单元格背景色"""
    shading_elm = OxmlElement("w:shd")
    shading_elm.set(qn("w:fill"), color)
    shading_elm.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_run_font(run, font_name="宋体", font_size=None, color=None, bold=None, italic=None):
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_pr.insert(0, r_fonts)
    if font_size is not None:
        run.font.size = Pt(font_size)
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def set_default_font(doc):
    style = doc.styles["Normal"]
    font = style.font
    font.name = "宋体"
    font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_centered_run(doc, text, font_size, font_name="宋体", color=None, bold=False):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_run_font(run, font_name=font_name, font_size=font_size, color=color, bold=bold)
    return paragraph


def add_cover(doc, title, subtitle=None, info_text=None, blank_lines=2):
    for _ in range(blank_lines):
        doc.add_paragraph()
    add_centered_run(doc, title, font_size=22, font_name="黑体", bold=True)
    if subtitle:
        add_centered_run(doc, f"—— {subtitle} ——", font_size=16, font_name="黑体", color=(0x44, 0x62, 0x88))
    if info_text:
        doc.add_paragraph()
        add_centered_run(doc, info_text, font_size=10, color=(0x99, 0x99, 0x99))
    doc.add_page_break()


def add_footer(doc):
    doc.add_paragraph()
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(FOOTER_TEXT)
    set_run_font(run, font_size=9, color=(0xAA, 0xAA, 0xAA), italic=True)


def extract_questions(data: dict) -> list | None:
    for key, value in data.items():
        if key != "info" and isinstance(value, list):
            return value
    return None


def load_assignment(json_path: str) -> dict:
    data = read_answer_json(json_path)
    info = data.get("info", {})
    questions = extract_questions(data)
    if not questions:
        raise ValueError(f"JSON文件中未找到题目数据: {json_path}")
    return {
        "json_path": json_path,
        "course_name": info.get("course_name", "未知课程"),
        "work_name": info.get("work_name", "未知作业"),
        "questions": questions,
    }


def iter_assignments(json_dir: str):
    for json_file in answer_json_files(json_dir):
        json_path = os.path.join(json_dir, json_file)
        try:
            yield load_assignment(json_path)
        except (OSError, JSONDecodeError, TypeError, ValueError) as e:
            print(f"导出失败 {json_file}: {e}")


def format_answer(answer) -> str:
    if isinstance(answer, list):
        return "；".join(str(item) for item in answer)
    return str(answer)


def option_letter(option: str) -> str:
    if "." in option:
        return option.split(".", 1)[0].strip()
    return option.split(" ", 1)[0].strip()


def is_correct_option(question_type: str, option: str, answer) -> bool:
    answer_text = format_answer(answer).strip()
    if question_type in ("单选题", "判断题"):
        return option.startswith(answer_text)
    if question_type == "多选题":
        return option_letter(option) in set(answer_text)
    return False


def add_options_table(doc, question_type: str, options: list, answer, font_size=10.5, color_answer=True):
    num_cols = 2
    num_rows = (len(options) + 1) // 2 or 1
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for index, option in enumerate(options):
        row = index // 2
        col = index % 2
        cell = table.cell(row, col)
        cell.paragraphs[0].clear()
        run = cell.paragraphs[0].add_run(option)
        set_run_font(run, font_size=font_size)

        if is_correct_option(question_type, option, answer):
            set_cell_shading(cell, "D4EDDA")
            run.bold = True
            if color_answer:
                run.font.color.rgb = RGBColor(0x15, 0x57, 0x24)


def add_question(doc, question: dict, index: int, compact=False, show_separator=False):
    question_type = question.get("type", "未知题型")
    title = question.get("title", "无题目")
    answer = question.get("answer", "无答案")
    options = question.get("option") or []
    label_color = TYPE_COLORS.get(question_type, TYPE_COLORS["其他"])

    header = doc.add_paragraph()
    header.paragraph_format.space_before = Pt(8 if compact else 16)
    header.paragraph_format.space_after = Pt(4 if compact else 6)

    number_text = f"{index}. " if compact else f"第{index}题 "
    number_run = header.add_run(number_text)
    set_run_font(number_run, font_name="微软雅黑", font_size=11 if compact else 13, bold=True)

    type_text = f"[{question_type}]" if compact else f" [{question_type}]"
    type_run = header.add_run(type_text)
    set_run_font(type_run, font_name="微软雅黑", font_size=9 if compact else 11, color=label_color, bold=True)

    text = doc.add_paragraph()
    text.paragraph_format.left_indent = Cm(0.5)
    text.paragraph_format.space_after = Pt(4 if compact else 8)
    set_run_font(text.add_run(title), font_size=10.5 if compact else 12)

    if options:
        add_options_table(doc, question_type, options, answer, font_size=9 if compact else 10.5, color_answer=not compact)
        if not compact:
            doc.add_paragraph()

    answer_paragraph = doc.add_paragraph()
    answer_paragraph.paragraph_format.left_indent = Cm(0.5)
    if not compact:
        answer_paragraph.paragraph_format.space_before = Pt(4)

    answer_label = answer_paragraph.add_run("答案：")
    set_run_font(answer_label, font_size=10 if compact else 11, color=ANSWER_COLOR, bold=True)

    answer_run = answer_paragraph.add_run(format_answer(answer))
    set_run_font(answer_run, font_size=10 if compact else 11, color=ANSWER_COLOR)

    if show_separator:
        separator = doc.add_paragraph()
        separator.paragraph_format.space_before = Pt(8)
        separator.paragraph_format.space_after = Pt(2)
        set_run_font(separator.add_run("─" * 50), font_size=8, color=(0xDD, 0xDD, 0xDD))


def add_assignment(doc, assignment: dict, index: int | None = None, compact=False):
    questions = assignment["questions"]
    if index is not None:
        title = f"{index}. {assignment['course_name']} — {assignment['work_name']}（共{len(questions)}题）"
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(12)
        paragraph.paragraph_format.space_after = Pt(6)
        set_run_font(paragraph.add_run(title), font_name="黑体", font_size=16 if not compact else 15, bold=True)

    for question_index, question in enumerate(questions, 1):
        add_question(
            doc,
            question,
            question_index,
            compact=compact,
            show_separator=not compact and question_index < len(questions),
        )


def new_document():
    doc = Document()
    set_default_font(doc)
    return doc


def generate_docx_from_json(json_path: str, output_path: str = None) -> str:
    """
    将单个作业的JSON文件导出为Word文档

    :param json_path: JSON答案文件的路径
    :param output_path: 输出Word文件的路径（可选，默认与JSON同目录同文件名）
    :return: 生成的docx文件路径
    """
    assignment = load_assignment(json_path)
    doc = new_document()
    add_cover(
        doc,
        assignment["course_name"],
        subtitle=assignment["work_name"],
        info_text=f"导出时间：{datetime.now().strftime('%Y年%m月%d日 %H:%M')}    题目总数：{len(assignment['questions'])} 题",
    )
    add_assignment(doc, assignment)
    add_footer(doc)

    if output_path is None:
        base_dir = os.path.dirname(json_path)
        base_name = os.path.splitext(os.path.basename(json_path))[0]
        output_path = os.path.join(base_dir, f"{base_name}.docx")

    doc.save(output_path)
    return output_path


def batch_export_to_docx(json_dir: str, output_dir: str = None, merge: bool = False) -> list:
    """
    批量将answers目录下的所有JSON文件导出为Word文档

    :param json_dir: JSON答案文件所在目录
    :param output_dir: 输出目录（可选，默认与json_dir相同）
    :param merge: 是否合并为一个文档（True=合并为一个文件，False=每个JSON单独导出）
    :return: 生成的docx文件路径列表
    """
    output_dir = output_dir or json_dir
    os.makedirs(output_dir, exist_ok=True)
    json_files = answer_json_files(json_dir)

    if not json_files:
        raise FileNotFoundError(f"目录中未找到JSON文件: {json_dir}")

    if not merge:
        generated_files = []
        for json_file in json_files:
            json_path = os.path.join(json_dir, json_file)
            output_path = os.path.join(output_dir, os.path.splitext(json_file)[0] + ".docx")
            try:
                generated_files.append(generate_docx_from_json(json_path, output_path))
            except (OSError, JSONDecodeError, TypeError, ValueError) as e:
                print(f"导出失败 {json_file}: {e}")
        return generated_files

    assignments = list(iter_assignments(json_dir))
    if not assignments:
        raise ValueError(f"目录中未找到可导出的JSON文件: {json_dir}")

    doc = new_document()
    add_cover(
        doc,
        "学习通作业汇总",
        info_text=f"导出时间：{datetime.now().strftime('%Y年%m月%d日 %H:%M')}    共 {len(assignments)} 份作业",
        blank_lines=3,
    )
    for index, assignment in enumerate(assignments, 1):
        add_assignment(doc, assignment, index=index, compact=True)
        if index < len(assignments):
            doc.add_page_break()
    add_footer(doc)

    output_path = os.path.join(output_dir, f"学习通作业汇总_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx")
    doc.save(output_path)
    return [output_path]


def export_course_assignments_to_docx(course_name: str, json_dir: str, output_path: str = None) -> str:
    """
    导出指定课程的所有作业到一个Word文档

    :param course_name: 课程名称（用于筛选JSON文件中的课程）
    :param json_dir: JSON答案文件目录
    :param output_path: 输出路径（可选）
    :return: 生成的docx文件路径
    """
    json_files = answer_json_files(json_dir)
    if not json_files:
        raise FileNotFoundError(f"目录中未找到JSON文件: {json_dir}")

    assignments = [
        assignment
        for assignment in iter_assignments(json_dir)
        if course_matches_query(assignment["course_name"], course_name)
    ]
    if not assignments:
        raise ValueError(f"未找到课程'{course_name}'的答案文件")

    title = assignments[0]["course_name"] if len({item["course_name"] for item in assignments}) == 1 else course_name
    doc = new_document()
    add_cover(
        doc,
        title,
        info_text=f"作业汇总  ·  共 {len(assignments)} 份作业  ·  {datetime.now().strftime('%Y年%m月%d日')}",
        blank_lines=3,
    )
    for index, assignment in enumerate(assignments, 1):
        add_assignment(doc, assignment, index=index, compact=True)
        if index < len(assignments):
            doc.add_page_break()
    add_footer(doc)

    if output_path is None:
        output_path = os.path.join(json_dir, f"{sanitize_filename(title)}_作业汇总.docx")

    doc.save(output_path)
    return output_path
